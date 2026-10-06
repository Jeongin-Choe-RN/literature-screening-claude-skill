"""Chat-only literature screening ledger. Local files only; no model or network calls."""
import base64
import copy
import csv
import hashlib
import io
import json
import re
import uuid
import zlib
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

SCHEMA = 'literature-screening.chat.v1'
POLICY = 'decision.all-required.v1'
DECISIONS = ('include', 'maybe', 'exclude')
STAGES = ('tiab', 'fulltext')
MAX_STATE = 64 * 1024 * 1024


def require(ok, message):
    if not ok:
        raise ValueError(message)


def norm(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def normalize_records(records, source_file=''):
    require(isinstance(records, list) and records, 'Provide actual records; the list is empty.')
    given = [norm(r.get('record_id')) for r in records if norm(r.get('record_id'))]
    require(len(given) == len(set(given)), 'Duplicate supplied record IDs. Resolve without dropping records.')
    used, result, number = set(given), [], 1
    for original in records:
        r = copy.deepcopy(original)
        require(norm(r.get('title')), 'Every record needs its actual title.')
        if not norm(r.get('record_id')):
            while f'R{number:06d}' in used:
                number += 1
            r['record_id'] = f'R{number:06d}'
            used.add(r['record_id'])
        for field in ('record_id', 'title', 'abstract', 'authors', 'year', 'doi', 'pmid'):
            r[field] = norm(r.get(field))
        r['doi'] = re.sub(r'^https?://(?:dx\.)?doi\.org/', '', r['doi'], flags=re.I)
        r['source_file'] = norm(r.get('source_file')) or source_file
        result.append(r)
    return result


def import_records(path):
    """Read UTF-8 CSV, RIS, NBIB or PubMed-text; never fetch missing abstracts."""
    path = Path(path)
    text = path.read_text(encoding='utf-8-sig')
    if re.search(r'^PMID\s*-', text, re.M):
        format_name = 'PubMed/NBIB'
    elif re.search(r'^TY\s*-', text, re.M):
        format_name = 'RIS'
    elif path.suffix.lower() == '.csv':
        format_name = 'CSV'
    else:
        raise ValueError('Use a UTF-8 CSV, NBIB/PubMed text, or RIS export. Do not relabel an unsupported file.')
    records = []
    if format_name == 'CSV':
        aliases = {'record_id': ['record_id', 'id', '문헌 id', '문헌id'],
                   'title': ['title', 'article title', '제목'], 'abstract': ['abstract', '초록'],
                   'authors': ['authors', 'author', '저자'], 'year': ['year', 'publication year', '연도'],
                   'doi': ['doi'], 'pmid': ['pmid'], 'source_file': ['source_file', '출처']}
        try:
            dialect = csv.Sniffer().sniff(text[:10000], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        require(reader.fieldnames and len(reader.fieldnames) == len(set(reader.fieldnames)), 'CSV headers are missing or duplicated.')
        for row in reader:
            require(None not in row, 'CSV has extra fields; check quoting and delimiter.')
            if not any(norm(v) for v in row.values()):
                continue
            lower = {norm(k).lower(): v for k, v in row.items()}
            records.append({key: next((lower[a] for a in names if a in lower), '') for key, names in aliases.items()})
    else:
        current, last = {}, None
        def flush():
            nonlocal current, last
            if current:
                def get(*tags):
                    return next((' '.join(current[t]) for t in tags if current.get(t)), '')
                title = get('TI', 'T1')
                if title:
                    ids = get('AID') + ' ' + get('LID')
                    match = re.search(r'(10\.\S+?)\s*\[doi\]', ids, re.I)
                    doi = get('DO') or (match.group(1) if match else '')
                    date = get('DP', 'PY', 'Y1')
                    year = re.search(r'\b(?:19|20)\d{2}\b', date)
                    records.append(dict(title=title, abstract=get('AB', 'N2'), authors=get('FAU', 'AU', 'A1'),
                                        year=year.group(0) if year else '', doi=doi, pmid=get('PMID')))
                else:
                    raise ValueError('An exported record has no title; do not silently skip it.')
            current, last = {}, None
        for line in text.splitlines():
            match = re.match(r'^([A-Z0-9]{2,4})\s*-\s?(.*)$', line)
            if match:
                tag, value = match.groups()
                if tag in ('PMID', 'TY') and current:
                    flush()
                if tag == 'ER':
                    flush()
                else:
                    current.setdefault(tag, []).append(value)
                    last = tag
            elif line.strip():
                require(last is not None and line[0].isspace(), 'Unrecognized export line. Check the source format.')
                current[last][-1] += ' ' + line.strip()
        flush()
    result = normalize_records(records, path.name)
    return result, {'format': format_name, 'imported': len(result),
                    'missing_abstracts': [r['record_id'] for r in result if not r['abstract']]}


def duplicate_candidates(records):
    candidates = []
    for i, left in enumerate(records):
        for right in records[i+1:]:
            reasons = [key for key in ('pmid', 'doi') if norm(left.get(key)) and norm(left.get(key)).lower() == norm(right.get(key)).lower()]
            if norm(left['title']).casefold() == norm(right['title']).casefold() and norm(left.get('year')) == norm(right.get('year')):
                reasons.append('title/year')
            if reasons:
                candidates.append({'record_ids': [left['record_id'], right['record_id']], 'matched_on': reasons})
    return candidates


def new_project(records, criteria, *, question, scope, search_log=None, language='ko', tutorial=False, title='Literature screening'):
    require(language in ('en', 'ko'), 'Supported output languages: en, ko.')
    require(norm(question) and norm(scope), 'Research question and actual screening scope are required.')
    state = dict(schema=SCHEMA, version='1.0', project_id=str(uuid.uuid4()), title=title, language=language,
                 protocol=dict(question=question, criteria=copy.deepcopy(criteria), scope=scope,
                               search_log=copy.deepcopy(search_log or []), tutorial=bool(tutorial)),
                 records=normalize_records(records), duplicates={}, frozen_sha256=None,
                 stages={'tiab': {}, 'fulltext': {}}, fulltext_started=False, study_links={},
                 history=[], created_at=now(), updated_at=now())
    return validate(state)


def frozen_data(s):
    return [s['protocol'], s['records'], s['duplicates']]


def active_ids(s):
    return [r['record_id'] for r in s['records'] if r['record_id'] not in s['duplicates']]


def blank_review():
    return dict(preparation='unknown', preparation_text='', student=None, ai=None, final=None,
                ai_issue='', material=dict(retrieval='pending', identity='pending', read_status='pending',
                                           filename='', note='', pages=[]))


def event(s, kind, *, stage=None, rid=None, user_text='', before=None, after=None):
    s['history'].append(dict(at=now(), kind=kind, stage=stage, record_id=rid, user_text=user_text,
                             before=copy.deepcopy(before), after=copy.deepcopy(after)))
    s['updated_at'] = now()


def confirm_duplicates(s, mapping, *, user_text):
    t = copy.deepcopy(validate(s))
    require(not t['frozen_sha256'] and norm(user_text), 'Confirm duplicates before starting screening.')
    t['duplicates'] = copy.deepcopy(mapping)
    event(t, 'duplicates_confirmed', user_text=user_text, after=mapping)
    return validate(t)


def start_screening(s, *, user_text):
    t = copy.deepcopy(validate(s))
    require(not t['frozen_sha256'] and norm(user_text), 'The user must confirm the criteria, scope and duplicate review once.')
    t['frozen_sha256'] = digest(frozen_data(t))
    t['stages']['tiab'] = {rid: blank_review() for rid in active_ids(t)}
    event(t, 'screening_started', user_text=user_text)
    return validate(t)


def aggregate(criteria, assessments):
    statuses = {a['criterion_id']: a['status'] for a in assessments}
    failed = [c['id'] for c in criteria if statuses[c['id']] == ('not_met' if c['kind'] == 'include' else 'met')]
    unclear = [c['id'] for c in criteria if statuses[c['id']] == 'unclear']
    return ('exclude', failed) if failed else ('maybe', unclear) if unclear else ('include', [])


def evidence_sources(s, stage, rid, row):
    if stage == 'fulltext':
        m = row['material']
        require(m['retrieval'] == 'retrieved' and m['identity'] == 'confirmed' and m['read_status'] == 'complete',
                'Confirm PDF identity and complete reading before full-text assessment.')
        return {str(p['page']): p['text'] for p in m['pages']}
    r = next(r for r in s['records'] if r['record_id'] == rid)
    return {key: str(r.get(key, '')) for key in ('title', 'abstract', 'year', 'authors', 'doi', 'pmid')}


def check_assessments(s, stage, rid, row, assessments):
    criteria = s['protocol']['criteria']
    require(len(assessments) == len(criteria) and {a.get('criterion_id') for a in assessments} == {c['id'] for c in criteria},
            'Assess every criterion exactly once.')
    sources = evidence_sources(s, stage, rid, row)
    for a in assessments:
        require(a.get('status') in ('met', 'not_met', 'unclear') and norm(a.get('explanation')), 'Invalid status or missing explanation.')
        if a['status'] == 'unclear':
            require(a.get('basis') == 'insufficient' and a.get('evidence') in ('', None) and a.get('location') in ('', None),
                    'Unclear: use insufficient basis, empty evidence and location; explain what is missing.')
        else:
            require(a.get('basis') in ('direct', 'counterevidence'), 'Evidence must be direct or counterevidence.')
            location = str(a.get('location', ''))
            require(location in sources and norm(a.get('evidence')) and norm(a['evidence']) in norm(sources[location]),
                    'Quote must occur verbatim in the stated source field or actual PDF page.')


def validate(s):
    require(s.get('schema') == SCHEMA and s.get('language') in ('ko', 'en'), 'Unsupported chat-only state. Keep older workbooks intact.')
    ids = [r['record_id'] for r in s['records']]
    require(ids and len(ids) == len(set(ids)) and all(norm(r.get('title')) for r in s['records']), 'Invalid record identities.')
    cs = s['protocol']['criteria']
    require(cs and any(c.get('kind') == 'include' for c in cs), 'At least one inclusion criterion is required.')
    require(len({c['id'] for c in cs}) == len(cs) and all(norm(c['id']) and norm(c['text']) and c['kind'] in ('include','exclude') for c in cs), 'Invalid criteria.')
    for removed, kept in s['duplicates'].items():
        require(removed in ids and kept in ids and removed != kept and kept not in s['duplicates'], 'Duplicates must map directly to a retained record.')
    require(set(s['study_links']).issubset(set(ids)) and all(norm(x) for x in s['study_links'].values()), 'Invalid study links.')
    if s['frozen_sha256']:
        require(s['frozen_sha256'] == digest(frozen_data(s)), 'Criteria, source records or duplicate decisions changed after screening began.')
        require(set(s['stages']['tiab']) == set(active_ids(s)), 'Missing or extra Stage 1 records.')
    else:
        require(not any(s['stages'].values()) and not s['fulltext_started'], 'Screening has not been confirmed.')
    if s['fulltext_started']:
        require(all((r['final'] or {}).get('decision') in ('include','exclude') for r in s['stages']['tiab'].values()), 'Stage 1 has unresolved decisions.')
        target = {rid for rid,r in s['stages']['tiab'].items() if r['final']['decision'] == 'include'}
        require(set(s['stages']['fulltext']) == target, 'Full-text targets do not match Stage 1 final inclusions.')
    else:
        require(not s['stages']['fulltext'], 'Start Stage 2 explicitly.')
    for stage in STAGES:
        for rid, row in s['stages'][stage].items():
            require(row['preparation'] in ('unknown','yes','waived'), 'Invalid private preparation state.')
            m = row['material']
            require(m['retrieval'] in ('pending','retrieved','not_retrieved') and m['read_status'] in ('pending','complete','partial','error') and m['identity'] in ('pending','confirmed'), 'Invalid PDF state.')
            pages=m['pages']
            require(all(isinstance(p.get('page'),int) and not isinstance(p['page'],bool) and p['page']>0 and isinstance(p.get('text'),str) for p in pages), 'Use actual 1-based PDF page numbers.')
            require(len({p['page'] for p in pages})==len(pages), 'Duplicate PDF page numbers.')
            if m['read_status']=='complete':
                require(m['retrieval']=='retrieved' and m['identity']=='confirmed' and any(norm(p['text']) for p in pages), 'Complete reading requires identified PDF text.')
            if m['retrieval']=='not_retrieved':
                require(norm(m['note']) and not row['ai'] and not row['final'], 'Missing full text is separate from eligibility decisions.')
            for role in ('student','ai','final'):
                d=row[role]
                if d:
                    require(d.get('decision') in DECISIONS, 'Unknown decision.')
                    if role!='ai':
                        require(norm(d.get('reason')) and norm(d.get('user_text')), 'Use actual student judgment and reason.')
                        if d['decision']=='exclude':
                            require(d.get('criterion_id') in {c['id'] for c in cs}, 'Link exclusions to an approved criterion.')
            if row['student']:
                require(row['student']['prewritten'] in ('yes','no','unknown'), 'Record whether the initial answer was prewritten.')
            if row['ai']:
                a=row['ai'];check_assessments(s,stage,rid,row,a['assessments'])
                require(a['decision']==aggregate(cs,a['assessments'])[0] and a['decision_policy']==POLICY, 'Invalid rule-derived AI judgment.')
                require(a['context']['student_seen'] in ('yes','no','unknown') and a['context']['previous_stage_seen'] in ('yes','no','unknown'), 'Record prior exposure honestly.')
            if row['final']:
                require(row['student'] and row['ai'] and row['final'].get('confirmed_by')=='student', 'A final comparison decision needs student initial, AI and explicit student confirmation.')
    return s


def update_row(s, stage, rid, kind, fn, user_text=''):
    t=copy.deepcopy(validate(s))
    require(stage in STAGES and rid in t['stages'][stage], 'Unknown stage or record ID.')
    row=t['stages'][stage][rid]
    before={k:copy.deepcopy(row[k]) for k in ('student','ai','final','ai_issue','preparation')}
    fn(row)
    after={k:copy.deepcopy(row[k]) for k in before}
    event(t,kind,stage=stage,rid=rid,user_text=user_text,before=before,after=after)
    return validate(t)


def prepare_batch(s, stage, ids, *, user_text, waive=False):
    require(norm(user_text) and ids and len(ids)==len(set(ids)), 'Use explicit confirmation for this actual batch.')
    t=copy.deepcopy(validate(s))
    for rid in ids:
        def apply(row):
            require(not row['ai'], 'This record already has an AI assessment.')
            row.update(preparation='waived' if waive else 'yes', preparation_text=user_text)
        t=update_row(t,stage,rid,'preparation',apply,user_text)
    return t


def set_material(s, rid, *, retrieval, filename='', pages=None, identity='pending', read_status='pending', note='', user_text=''):
    def apply(row):
        require(not row['ai'] and not row['final'], 'Do not replace evaluated source material. Start a documented new round.')
        require(retrieval!='not_retrieved' or (norm(user_text) and norm(note)), 'Confirm that the PDF actually could not be obtained.')
        row['material']=dict(retrieval=retrieval,filename=filename,pages=copy.deepcopy(pages or []),identity=identity,read_status=read_status,note=note)
    return update_row(s,'fulltext',rid,'material',apply,user_text)


def record_ai(s, stage, rid, assessments, *, student_seen='unknown', previous_stage_seen='unknown', model=None, context_note=''):
    def apply(row):
        require(not row['ai'] and row['preparation'] in ('yes','waived'), 'Confirm private preparation or explicit waiver before revealing AI; do not overwrite saved AI.')
        check_assessments(s,stage,rid,row,assessments)
        value, criterion_ids=aggregate(s['protocol']['criteria'],assessments)
        row['ai']=dict(decision=value,criterion_ids=criterion_ids,assessments=copy.deepcopy(assessments),decision_policy=POLICY,at=now(),model=model,
                       context=dict(student_seen='yes' if row['student'] else student_seen,previous_stage_seen=previous_stage_seen,
                                    verification='not_verified',note=context_note,preparation=row['preparation']))
        row['ai_issue']=''
    return update_row(s,stage,rid,'ai',apply)


def record_ai_issue(s, stage, rid, note):
    def apply(row):
        require(not row['ai'] and norm(note), 'Keep successful AI results; record a concrete error only.')
        row['ai_issue']=note
    return update_row(s,stage,rid,'ai_issue',apply)


def record_student(s, stage, rid, decision, reason, *, user_text, prewritten='unknown', criterion_id=None):
    def apply(row):
        require(not row['student'], 'Preserve the original student judgment; use a new documented round for corrections.')
        row['student']=dict(decision=decision,reason=reason,criterion_id=criterion_id,user_text=user_text,prewritten=prewritten,at=now())
    return update_row(s,stage,rid,'student_initial',apply,user_text)


def record_final(s, stage, rid, decision, reason, *, user_text, criterion_id=None):
    require(stage!='tiab' or not s['fulltext_started'], 'Stage 1 decisions are locked once Stage 2 starts.')
    def apply(row):
        row['final']=dict(decision=decision,reason=reason,criterion_id=criterion_id,user_text=user_text,confirmed_by='student',at=now())
    return update_row(s,stage,rid,'student_final',apply,user_text)


def confirm_matching(s, stage, ids, *, user_text):
    require(ids and len(ids)==len(set(ids)) and norm(user_text), 'Confirm the actual listed IDs and their evidence.')
    t=copy.deepcopy(validate(s))
    for rid in ids:
        row=t['stages'][stage][rid]
        require(row['student'] and row['ai'] and not row['final'] and row['student']['decision']==row['ai']['decision'] and row['ai']['decision']!='maybe', 'Only unfinalized include/include or exclude/exclude matches can be confirmed together.')
        d=row['student']
        t=record_final(t,stage,rid,d['decision'],d['reason'],criterion_id=d.get('criterion_id'),user_text=user_text)
    return t


def start_fulltext(s, *, user_text):
    t=copy.deepcopy(validate(s))
    require(t['frozen_sha256'] and not t['fulltext_started'] and norm(user_text), 'Confirm transition once; do not restart completed Stage 2 work.')
    require(all((r['final'] or {}).get('decision') in ('include','exclude') for r in t['stages']['tiab'].values()), 'Resolve Stage 1 holds and pending decisions first.')
    t['stages']['fulltext']={rid:blank_review() for rid,r in t['stages']['tiab'].items() if r['final']['decision']=='include'}
    t['fulltext_started']=True
    event(t,'fulltext_started',user_text=user_text)
    return validate(t)


def link_studies(s, mapping, *, user_text):
    t=copy.deepcopy(validate(s));require(norm(user_text),'Confirm report-to-study links before counting distinct studies.')
    t['study_links'].update(mapping);event(t,'study_links',user_text=user_text,after=mapping)
    return validate(t)


def comparison(s, stage):
    rows=list(s['stages'][stage].values())
    eligible=[r for r in rows if stage=='tiab' or r['material']['retrieval']=='retrieved']
    pairs=[r for r in eligible if r['student'] and r['ai']]
    same=[r for r in pairs if r['student']['decision']==r['ai']['decision']]
    return dict(targets=len(rows),denominator=len(pairs),pending=len(eligible)-len(pairs),agreement=len(same),
                definite_agreements=sum(r['ai']['decision']!='maybe' for r in same),hold_agreements=sum(r['ai']['decision']=='maybe' for r in same),
                percent=100*len(same)/len(pairs) if pairs else None,
                exposure_counts=dict(Counter(r['ai']['context']['student_seen'] for r in pairs)))


def summary(s):
    validate(s)
    ta=s['stages']['tiab'];ft=s['stages']['fulltext']
    ta_final=Counter((r['final'] or {}).get('decision','unset') for r in ta.values())
    ta_pending=sum(v for k,v in ta_final.items() if k not in ('include','exclude'))
    if not s['frozen_sha256']:
        ta_pending=len(active_ids(s))
    retrieved=[r for r in ft.values() if r['material']['retrieval']=='retrieved']
    missing=sum(r['material']['retrieval']=='not_retrieved' for r in ft.values())
    retrieval_pending=sum(r['material']['retrieval']=='pending' for r in ft.values())
    ft_final=Counter((r['final'] or {}).get('decision','unset') for r in retrieved)
    pending=sum(v for k,v in ft_final.items() if k not in ('include','exclude'))
    reasons=Counter(r['final']['criterion_id'] for r in retrieved if (r['final'] or {}).get('decision')=='exclude')
    included=[rid for rid,r in ft.items() if (r['final'] or {}).get('decision')=='include']
    complete=bool(s['frozen_sha256']) and ta_pending==0 and (ta_final['include']==0 or (s['fulltext_started'] and retrieval_pending==0 and pending==0))
    require(len(ft)==missing+retrieval_pending+len(retrieved),'Full-text retrieval counts do not reconcile.')
    require(len(retrieved)==ft_final['include']+ft_final['exclude']+pending,'Full-text eligibility counts do not reconcile.')
    require(sum(reasons.values())==ft_final['exclude'],'Exclusion reason counts do not reconcile.')
    studies=len({s['study_links'][rid] for rid in included}) if complete and all(rid in s['study_links'] for rid in included) else None
    return dict(complete=complete,records_imported=len(s['records']),duplicates_removed=len(s['duplicates']),
                records_screening_scope=len(active_ids(s)),records_with_initial=sum(bool(r['student']) for r in ta.values()),
                records_excluded=ta_final['exclude'],records_to_fulltext=ta_final['include'],records_pending_final=ta_pending,
                fulltext_started=s['fulltext_started'],reports_sought=len(ft),reports_not_retrieved=missing,
                reports_retrieval_pending=retrieval_pending,reports_retrieved=len(retrieved),
                reports_assessment_started=sum(bool(r['student'] or r['ai']) for r in retrieved),
                reports_excluded=ft_final['exclude'],reports_included=ft_final['include'],reports_pending_final=pending,
                fulltext_exclusion_reasons=dict(reasons),studies_included=studies,
                comparison={stage:comparison(s,stage) for stage in STAGES})


def translated(s, english, korean):
    return korean if s['language']=='ko' else english


def label(s, value):
    return ({'include':'포함','maybe':'보류','exclude':'제외'} if s['language']=='ko' else
            {'include':'Include','maybe':'Hold','exclude':'Exclude'}).get(value, '')


def tables(s):
    """All visible tables are derived from the ledger; no decision cells contain formulas."""
    counts=summary(s)
    tr=lambda en,ko:translated(s,en,ko)
    result={}
    def add(en,ko,en_headers,ko_headers,rows):
        result[tr(en,ko)]=(ko_headers if s['language']=='ko' else en_headers,list(rows))
    add('Summary','안내·요약',['Item','Value'],['항목','내용'],[
        [tr('Version','버전'),'Chat 1.0'],[tr('Project','작업'),s['title']],
        [tr('Research question','연구질문'),s['protocol']['question']],
        [tr('Scope','선별 범위'),s['protocol']['scope']],
        [tr('Material','자료'),tr('Fictional practice — not research data','가상 연습 · 연구자료 아님') if s['protocol']['tutorial'] else tr('User-provided records','사용자가 제공한 문헌')],
        [tr('Status','상태'),tr('Complete within the declared scope','선언한 범위 내 기록 완료') if counts['complete'] else tr('In progress','진행 중')],
        [tr('Workflow','진행 순서'),tr('Read and privately record → AI assessment → submit your initial decisions → compare → confirm final decisions','먼저 읽고 따로 기록 → AI 평가 → 내 최초 판정 제출 → 비교 → 최종 확인')],
        [tr('Resume','이어 하기'),tr('Attach the latest workbook in chat. Reattach required PDFs. Direct edits require confirmation.','최신 엑셀을 대화에 첨부합니다. 필요한 PDF를 다시 첨부하고 직접 수정한 값을 확인합니다.')],
        [tr('Agreement','일치율 해석'),tr('Agreement is not accuracy or verified independence. Stage-specific denominator: records with both judgments. Hold–hold matches are separate.','일치율은 정확도·검증된 독립성 지표가 아닙니다. 단계별로 두 판정이 있는 문헌을 분모로 삼고 보류끼리 일치를 별도 표시합니다.')],
        [tr('Saved at','저장 시각'),s['updated_at']]])
    add('Criteria','선별기준',['Criterion ID','Kind','Criterion'],['기준 ID','종류','기준'],
        [[c['id'],tr(c['kind'],'포함' if c['kind']=='include' else '제외'),c['text']] for c in s['protocol']['criteria']])
    add('Records','문헌목록',['Record ID','Title','Abstract','Authors','Year','DOI','PMID','Source','Duplicate of','Study ID'],
        ['문헌 ID','제목','초록','저자','연도','DOI','PMID','출처','중복 원본 ID','연구 ID'],
        [[r['record_id'],r['title'],r.get('abstract'),r.get('authors'),r.get('year'),r.get('doi'),r.get('pmid'),r.get('source_file'),s['duplicates'].get(r['record_id']),s['study_links'].get(r['record_id'])] for r in s['records']])
    records={r['record_id']:r for r in s['records']}
    headers=['Record ID','Title','Student initial','AI decision','Initial match','Student final','Primary exclusion criterion','Final reason','Student initial reason','Prewritten','AI saw student answer','AI issue','Retrieval','Reading','PDF/source','Material note']
    ko_headers=['문헌 ID','제목','학생 최초 판정','AI 판정','초기 일치','학생 최종 판정','주된 제외 기준','최종 근거','학생 최초 근거','사전 기록 여부','AI의 학생 답 사전 열람','AI 미완료 사항','원문 확보','읽기 상태','PDF·출처','자료 확인 사항']
    evidence,runs=[],[]
    for stage in STAGES:
        rows=[]
        for rid,row in s['stages'][stage].items():
            h,a,f=row['student'] or {},row['ai'] or {},row['final'] or {}
            m=row['material']
            match=tr('Same','일치') if h and a and h['decision']==a['decision'] else tr('Different','불일치') if h and a else tr('Pending','비교 전')
            rows.append([rid,records[rid]['title'],label(s,h.get('decision')),label(s,a.get('decision')),match,
                         label(s,f.get('decision')),f.get('criterion_id') if f.get('decision')=='exclude' else None,
                         f.get('reason'),h.get('reason'),h.get('prewritten'),a.get('context',{}).get('student_seen'),row['ai_issue'],
                         m['retrieval'] if stage=='fulltext' else None,m['read_status'] if stage=='fulltext' else None,m['filename'],m['note']])
            for item in a.get('assessments',[]):
                evidence.append([f'{stage}:{rid}:{item["criterion_id"]}',stage,rid,item['criterion_id'],item['status'],item.get('evidence'),item.get('location'),item['basis'],item['explanation']])
            if a:
                runs.append([f'{stage}:{rid}',stage,rid,a['model'],a['at'],POLICY,a['context']['student_seen'],a['context']['previous_stage_seen'],'not_verified',a['context']['preparation'],a['context']['note']])
        add('Stage 1' if stage=='tiab' else 'Stage 2','1차선별' if stage=='tiab' else '2차선별',headers,ko_headers,rows)
    add('AI Evidence','AI기준별근거',['Evidence ID','Stage','Record ID','Criterion ID','Status','Verbatim evidence','Source field / PDF page','Basis','Explanation'],
        ['근거 ID','단계','문헌 ID','기준 ID','평가','원문 인용','출처 필드·PDF 쪽','근거 유형','설명'],evidence)
    add('AI Runs','AI평가기록',['Evaluation ID','Stage','Record ID','Model','Time','Decision rule','Student seen','Previous stage seen','Independence','Preparation','Context note'],
        ['평가 ID','단계','문헌 ID','모델','시각','종합 규칙','학생 답 사전 열람','이전 단계 열람','독립성 검증','사전 준비','상황 기록'],runs)
    rows=[]
    for stage in STAGES:
        c=counts['comparison'][stage]
        rows += [[stage,tr('Target records','단계 대상'),c['targets']],
                 [stage,tr('Comparison denominator','비교 분모'),c['denominator']],
                 [stage,tr('Awaiting comparison among available records','현재 확보 자료 중 비교 대기'),c['pending']],
                 [stage,tr('Matching Include/Exclude decisions','포함·제외끼리 일치'),c['definite_agreements']],
                 [stage,tr('Hold–Hold matches','보류끼리 일치'),c['hold_agreements']],
                 [stage,tr('Agreement (%)','일치율 (%)'),c['percent']],
                 [stage,tr('Exposure reports: yes / no / unknown','사전 열람 신고: yes / no / unknown'),c['exposure_counts']]]
    add('Comparison','판정비교',['Entry','Stage','Measure','Value'],['항목 ID','단계','지표','값'],[[str(i+1)]+r for i,r in enumerate(rows)])
    fields=[('Records imported','불러온 문헌','records_imported'),('Duplicates removed','중복 제거','duplicates_removed'),
            ('Records in screening scope','1차 대상','records_screening_scope'),('Student initial recorded','학생 최초 판정 기록','records_with_initial'),
            ('Stage 1 excluded','1차 제외','records_excluded'),('Stage 1 final pending','1차 최종 판단 대기','records_pending_final'),
            ('Reports selected for retrieval','원문 검토 대상으로 선정','records_to_fulltext'),('Reports sought: Stage 2 started','원문 확보 대상: 2차 착수','reports_sought'),
            ('Reports not retrieved','원문 미확보','reports_not_retrieved'),('Retrieval pending','원문 확보 확인 중','reports_retrieval_pending'),
            ('Reports retrieved','확보 원문','reports_retrieved'),('Full-text assessment started','원문 평가 착수','reports_assessment_started'),
            ('Full-text excluded','원문 제외','reports_excluded'),('Full-text final pending','확보 원문 최종 판단 대기','reports_pending_final'),
            ('Reports included','포함 보고서','reports_included'),('Studies included','포함 연구','studies_included')]
    prisma=[[tr(en,ko),counts[key]] for en,ko,key in fields]
    prisma += [[tr('Full-text exclusion criterion: ','원문 제외 기준: ')+cid,n] for cid,n in counts['fulltext_exclusion_reasons'].items()]
    prisma += [[tr('Scope','범위'),s['protocol']['scope']],
               [tr('Interpretation','해석'),tr('Counts for a later diagram, not a completed PRISMA diagram. Unknown is blank, actual zero is 0.','흐름도 작성용 건수이며 완성된 그림이 아닙니다. 미확인은 빈칸, 실제 0은 숫자 0입니다.')]]
    add('PRISMA Counts','PRISMA집계',['Measure','Value'],['항목','값'],prisma)
    add('Search Log','검색기록',['Entry','Original search record'],['기록 ID','원본 검색 기록'],[[str(i+1),r] for i,r in enumerate(s['protocol']['search_log'])])
    add('History','변경기록',['Event','Time','Type','Stage','Record ID','Student statement','Before','After'],
        ['기록 ID','시각','유형','단계','문헌 ID','학생 입력','변경 전','변경 후'],
        [[str(i+1),e['at'],e['kind'],e['stage'],e['record_id'],e['user_text'],e['before'],e['after']] for i,e in enumerate(s['history'])])
    return result


def cell_value(value):
    if isinstance(value,(dict,list)):
        value=json.dumps(value,ensure_ascii=False,sort_keys=True)
    if isinstance(value,str):
        # Source and audit text remain complete in the resume ledger.
        value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '\uFFFD', value)
        if len(value)>32000:
            value=value[:31800]+'\n[Long display shortened; complete text is preserved in the resume ledger.]'
    return value


def add_sheet(wb,name,headers,rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    ws=wb.create_sheet(name)
    for n,values in enumerate([headers]+list(rows),1):
        for col,value in enumerate(values,1):
            cell=ws.cell(n,col);value=cell_value(value)
            cell.value=value
            if isinstance(value,str):
                cell.data_type='s'  # Imported titles starting with '=' are text, never formulas.
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245EB7')
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=85
    for i,h in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(i)].width=48 if any(t in h.lower() for t in ('title','reason','evidence','abstract','value','statement','record','제목','초록','근거','내용','기록')) else 23
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value in ('Include','포함','Exclude','제외','Hold','보류'):
                color='E5EFFC' if c.value in ('Include','포함') else 'FBECEB' if c.value in ('Exclude','제외') else 'FFF5DD'
                c.fill=PatternFill('solid',fgColor=color)
    return ws


def write_workbook(s,path):
    from openpyxl import Workbook
    validate(s)
    wb=Workbook();wb.remove(wb.active)
    for name,(headers,rows) in tables(s).items():
        add_sheet(wb,name,headers,rows)
    raw=json.dumps(s,ensure_ascii=False,separators=(',',':')).encode()
    require(len(raw)<=MAX_STATE,'Ledger exceeds the supported size. Preserve the state and split source-text storage explicitly.')
    encoded=base64.b64encode(zlib.compress(raw)).decode()
    resume=add_sheet(wb,'_resume',['Schema','SHA256','Encoding'],[[SCHEMA,hashlib.sha256(raw).hexdigest(),'zlib+base64']]+[[encoded[i:i+24000]] for i in range(0,len(encoded),24000)])
    resume.sheet_state='hidden';wb.save(path)
    restored,edits=read_workbook(path)
    require(restored==s and not edits,'Workbook did not round-trip faithfully.')
    return summary(s)


def read_workbook(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    require('_resume' in wb.sheetnames,'No chat-only resume ledger. Preserve this file and identify its original format.')
    ws=wb['_resume'];require(ws['A2'].value==SCHEMA,'This workbook belongs to another screening edition. Do not overwrite it.')
    encoded=''.join(str(ws.cell(i,1).value or '') for i in range(3,ws.max_row+1))
    require(len(encoded)<MAX_STATE*2,'Unexpectedly large resume payload.')
    data=base64.b64decode(encoded,validate=True)
    dec=zlib.decompressobj();raw=dec.decompress(data,MAX_STATE+1)
    require(len(raw)<=MAX_STATE and dec.eof and not dec.unused_data,'Invalid or oversized compressed ledger.')
    require(hashlib.sha256(raw).hexdigest()==ws['B2'].value,'Damaged resume ledger.')
    s=validate(json.loads(raw));edits=[]
    expected=tables(s)
    require(set(wb.sheetnames)==set(expected)|{'_resume'},'Worksheet added, removed or renamed. Preserve edits and reconcile manually.')
    for name,(headers,rows) in expected.items():
        visible=wb[name]
        actual_headers=[c.value for c in visible[1]]
        require(len(actual_headers)==len(headers) and set(actual_headers)==set(headers),f'Headers changed in {name}; reconcile before resuming.')
        actual={}
        for cells in visible.iter_rows(min_row=2):
            require(not any(c.data_type=='f' for c in cells),f'Formula in {name}. Confirm literal values before importing.')
            values=dict(zip(actual_headers,[c.value for c in cells]))
            if not any(v is not None for v in values.values()):
                continue
            key=values[headers[0]]
            require(key is not None and key not in actual,f'Missing or duplicated row ID in {name}.')
            actual[key]=values
        wanted={cell_value(r[0]):dict(zip(headers,[cell_value(v) for v in r])) for r in rows}
        require(set(actual)==set(wanted),f'Rows or identifiers changed in {name}. Preserve changes and reconcile before resuming.')
        for key,old in wanted.items():
            for header in headers:
                a,b=old.get(header),actual[key].get(header)
                if ('' if a is None else a)!=('' if b is None else b):
                    edits.append(dict(sheet=name,row_id=key,column=header,saved=a,edited=b))
    return s,edits


def write_student_form(s,stage,ids,path):
    """Optional offline form, deliberately separate from AI and saved judgments."""
    from openpyxl import Workbook
    validate(s);require(ids and len(ids)==len(set(ids)) and all(rid in s['stages'][stage] for rid in ids),'Unknown batch IDs.')
    tr=lambda en,ko:translated(s,en,ko)
    records={r['record_id']:r for r in s['records']}
    wb=Workbook();wb.remove(wb.active)
    add_sheet(wb,tr('Instructions','안내'),[tr('Item','항목'),tr('Value','내용')],[
        [tr('Use','사용'),tr('Read first and record your own decisions. Send this file only after the AI assessment. It is not a resume workbook.','먼저 읽고 자신의 판정을 기록합니다. AI 평가 뒤에만 보내세요. 진행 재개용 엑셀이 아닙니다.')],
        [tr('Stage','단계'),stage],[tr('Scope','범위'),s['protocol']['scope']]])
    headers=['Record ID','Title','Abstract','My initial decision','My reason','Exclusion criterion','PDF page']
    if s['language']=='ko':
        headers=['문헌 ID','제목','초록','내 최초 판정','내 근거','제외 기준','PDF 쪽']
    add_sheet(wb,tr('My Decisions','내 판정'),headers,[[rid,records[rid]['title'],records[rid].get('abstract') if stage=='tiab' else '',None,None,None,None] for rid in ids])
    add_sheet(wb,tr('Criteria','선별기준'),['ID',tr('Kind','종류'),tr('Criterion','기준')],[[c['id'],c['kind'],c['text']] for c in s['protocol']['criteria']])
    wb.save(path)


def save_state(s,path):
    validate(s)
    Path(path).write_text(json.dumps(s,ensure_ascii=False,indent=2),encoding='utf-8')


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser(description=__doc__)
    sub=parser.add_subparsers(dest='command',required=True)
    imp=sub.add_parser('import');imp.add_argument('source');imp.add_argument('records_json')
    out=sub.add_parser('workbook');out.add_argument('state_json');out.add_argument('output_xlsx')
    check=sub.add_parser('inspect');check.add_argument('workbook')
    args=parser.parse_args()
    if args.command=='import':
        records,info=import_records(args.source)
        Path(args.records_json).write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(info,ensure_ascii=False))
    elif args.command=='workbook':
        state=json.loads(Path(args.state_json).read_text(encoding='utf-8'))
        print(json.dumps(write_workbook(state,args.output_xlsx),ensure_ascii=False))
    else:
        state,edits=read_workbook(args.workbook)
        print(json.dumps({'summary':summary(state),'manual_edits':edits},ensure_ascii=False))
